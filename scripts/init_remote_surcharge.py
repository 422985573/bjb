# -*- coding: utf-8 -*-
"""从两份 xlsx 灌库：新大货文章 border / toll 快递报价表的偏远附加费。

数据源（config 可覆盖路径）：
  - border偏远费.xlsx：
      左表 A-F 列 = Suburb / State / Postcode / Zones / RAS Tier / BEX Delivery
      右表 K-N 列（约 5-12 行）= 「偏远地区附加费 - 包裹/大宗货物 - RASx 级」→ 金额
  - TOLL偏远费.xlsx：
      6 个列块（列 1/6/11/16/21/26），每格文本 = "<postcode> <suburb> <STATE> <fee>"

写入三张表（DELETE 后 INSERT，单事务，幂等，可重复运行）：
  border_ras_fee(ras_tier, parcel_fee, bulk_fee)
  border_remote_postcode(postcode, suburb, state, zone, ras_tier)
  toll_remote_postcode(postcode, suburb, state, fee)

运行方式: ./bjb_venv/bin/python scripts/init_remote_surcharge.py
"""
import os
import re
import sys

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE_DIR)

import openpyxl  # noqa: E402

import config  # noqa: E402
import db  # noqa: E402

# TOLL 每格文本：<postcode> <suburb...> <STATE> <fee>
_TOLL_CELL_RE = re.compile(r'^(\d{3,4})\s+(.*?)\s+([A-Z]{2,3})\s+(\d+(?:\.\d+)?)$')
# border 右表描述：偏远地区附加费 - 包裹/大宗货物 - RASx 级
_BORDER_FEE_RE = re.compile(r'(包裹|大宗货物).*?(RAS\s*[1-4])')
# TOLL 数据所在列（1-based）：每个块首列
_TOLL_DATA_COLS = [1, 6, 11, 16, 21, 26]


def _norm_pc(raw):
    """规范邮编为 4 位字符串；非 3/4 位数字返回 None。"""
    digits = ''.join(c for c in str(raw or '') if c.isdigit())
    if len(digits) in (3, 4):
        return digits.zfill(4)
    return None


def parse_border(path):
    """返回 (fee_rows, postcode_rows)。

    fee_rows: [(ras_tier, parcel_fee, bulk_fee), ...] 4 条
    postcode_rows: [(postcode, suburb, state, zone, ras_tier), ...]
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    # 右表：RAS 等级 -> 包裹费 / 大宗费（同一 tier 两行合并到一条）
    parcel = {}
    bulk = {}
    for r in range(1, ws.max_row + 1):
        desc = ws.cell(r, 11).value  # K 列描述
        amount = ws.cell(r, 13).value  # M 列金额
        if desc is None or amount is None:
            continue
        m = _BORDER_FEE_RE.search(str(desc))
        if not m:
            continue
        kind = m.group(1)
        tier = m.group(2).replace(' ', '')
        try:
            fee = float(amount)
        except (TypeError, ValueError):
            continue
        (parcel if kind == '包裹' else bulk)[tier] = fee
    fee_rows = []
    for tier in sorted(set(parcel) | set(bulk)):
        fee_rows.append((tier, parcel.get(tier, 0), bulk.get(tier, 0)))

    # 左表：逐行 postcode -> RAS tier
    postcode_rows = []
    for r in range(2, ws.max_row + 1):
        suburb = ws.cell(r, 1).value
        state = ws.cell(r, 2).value
        postcode = ws.cell(r, 3).value
        zone = ws.cell(r, 4).value
        tier = ws.cell(r, 5).value
        pc = _norm_pc(postcode)
        if pc is None or tier is None or not str(tier).strip():
            continue
        postcode_rows.append((
            pc,
            str(suburb or '').strip(),
            str(state or '').strip(),
            str(zone or '').strip(),
            str(tier).replace(' ', '').strip(),
        ))
    wb.close()
    return fee_rows, postcode_rows


def parse_toll(path):
    """返回 [(postcode, suburb, state, fee), ...]。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for r in range(1, ws.max_row + 1):
        for c in _TOLL_DATA_COLS:
            v = ws.cell(r, c).value
            if v is None:
                continue
            s = str(v).strip()
            if not s or s.startswith('Postcode'):  # 跳过表头
                continue
            m = _TOLL_CELL_RE.match(s)
            if not m:
                continue
            pc = _norm_pc(m.group(1))
            if pc is None:
                continue
            try:
                fee = float(m.group(4))
            except ValueError:
                continue
            rows.append((pc, m.group(2).strip(), m.group(3).strip(), fee))
    wb.close()
    return rows


def main():
    print('Step 0: ensure DB schema...')
    db.init_db()

    border_path = config.BORDER_REMOTE_XLSX
    toll_path = config.TOLL_REMOTE_XLSX
    for p in (border_path, toll_path):
        if not os.path.isfile(p):
            print(f'ERROR: xlsx not found: {p}')
            sys.exit(1)

    print('Step 1: parse border xlsx...')
    fee_rows, border_pc_rows = parse_border(border_path)
    print(f'  border: {len(fee_rows)} fee tiers, {len(border_pc_rows)} postcode rows')
    for fr in fee_rows:
        print(f'    {fr[0]}: 包裹 {fr[1]} / 大宗 {fr[2]}')

    print('Step 2: parse TOLL xlsx...')
    toll_rows = parse_toll(toll_path)
    print(f'  toll: {len(toll_rows)} postcode rows')

    print('Step 3: write to DB (single transaction)...')
    with db.get_db() as conn:
        cur = conn.cursor()
        cur.execute('DELETE FROM border_ras_fee')
        cur.executemany(
            'INSERT INTO border_ras_fee (ras_tier, parcel_fee, bulk_fee) VALUES (?, ?, ?)',
            fee_rows,
        )
        cur.execute('DELETE FROM border_remote_postcode')
        cur.executemany(
            'INSERT INTO border_remote_postcode (postcode, suburb, state, zone, ras_tier) '
            'VALUES (?, ?, ?, ?, ?)',
            border_pc_rows,
        )
        cur.execute('DELETE FROM toll_remote_postcode')
        cur.executemany(
            'INSERT INTO toll_remote_postcode (postcode, suburb, state, fee) VALUES (?, ?, ?, ?)',
            toll_rows,
        )
        conn.commit()
    print('  done.')
    print('\nAll done!')


if __name__ == '__main__':
    main()
